import re
from typing import Dict, List, Tuple
from docx import Document

PLACEHOLDER_PATTERNS = [
    r"\[[^\]]+\]",     # [PLACEHOLDER]
    r"\{\{[^}]+\}\}",  # {{PLACEHOLDER}}
    r"<[^>]+>",        # <PLACEHOLDER>
    r"\$\s*\[[^\]]+\]"
]

def _tokenize_words_with_offsets(text: str) -> List[Tuple[str, int]]:
    """Return list of (word, start_char_idx) so we can slice by word windows robustly."""
    out = []
    i = 0
    for m in re.finditer(r"\S+", text):
        out.append((m.group(0), m.start()))
    return out

def extract_placeholders(file_path: str, context_window_words: int = 80) -> Dict:
    doc = Document(file_path)

    # Combine paragraphs; skip empty lines to reduce noise
    lines: List[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(lines)

    # Find all matches with positions so we can build per-occurrence context
    # Find all matches with positions, avoiding duplicates for $[...] placeholders
    matches = []
    for pat in PLACEHOLDER_PATTERNS:
        for m in re.finditer(pat, full_text):
            raw = m.group(0)
            # ⚠️ Skip redundant [____] match if it's part of a $[____] already
            if pat == r"\[[^\]]+\]" and re.match(r"^\$\s*\[", full_text[m.start()-1 : m.end()]):
                continue
            matches.append((m.start(), m.end(), raw, pat))


    print("\n=== DEBUG: PLACEHOLDER TEST ===")
    print("Document length:", len(full_text))
    print("Patterns:", PLACEHOLDER_PATTERNS)
    print("Matches found:", len(matches))
    for s, e, raw, pat in matches:
        print(f"Matched [{raw}] using {pat}")
    print("==============================\n")

    # Keep reading order
    matches.sort(key=lambda x: x[0])

    # Tokenize words to extract context windows
    words_with_offs = _tokenize_words_with_offsets(full_text)
    word_positions = [off for _, off in words_with_offs]
    words_only = [w for w, _ in words_with_offs]

    occurrences = []
    context_map: Dict[str, str] = {}

    def normalize_label(raw: str) -> str:
        label = raw.strip().strip("[]{}<> ").upper()
        return label

    for idx, (s, e, raw, pat) in enumerate(matches):
        occ_id = str(idx)

        # Label normalization
        if pat == r"\$\s*\[[^\]]+\]":
            label = "$[__________]"  # special case for monetary placeholders
        else:
            label = normalize_label(raw)

        # Build local context window around the placeholder
        left_idx = 0
        right_idx = len(words_only) - 1
        for i, off in enumerate(word_positions):
            if off <= s:
                left_idx = i
            if off < e:
                right_idx = i

        cstart = max(0, left_idx - context_window_words)
        cend = min(len(words_only), right_idx + 1 + context_window_words)
        snippet = " ".join(words_only[cstart:cend]).strip()

        occurrences.append({"id": occ_id, "label": label})
        context_map[occ_id] = snippet

    # ✅ Clean ordered return
    return {
        "text_preview": full_text[:500] + ("..." if len(full_text) > 500 else ""),
        "occurrences": occurrences,  
        "context_map": context_map  
    }
